"""
Seed villages + health workers from the roster Excel
("list of health centres of 2 blocks of Raipur district...") into the
villages and health_workers tables, linking villages to their SHC in
health_facilities.

Requires seed_facilities.py to have been run first (SHCs must already
exist in health_facilities).

IDEMPOTENT: safe to re-run. Villages are upserted by `code` (or by
name+block if a row has no code). Workers are upserted by
(facility_id, role) — a re-run updates name/mobile/status instead of
duplicating rows, so admin edits made through a future admin panel are
the only source you'd overwrite, and only by re-running this script.

Run once from backend folder:
    python seed_villages.py
Note: the roster xlsx and this script should be in the same folder.
"""

import asyncio
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import load_workbook
from sqlalchemy import select
from app.core.database import AsyncSessionLocal
from app.models.models import HealthFacility, Village, HealthWorker, WorkerRole, WorkerStatus

ROSTER_FILE = os.path.join(
    os.path.dirname(__file__),
    "list of health centres of 2 blocks of Raipur district with name of health care providers.xlsx",
)

# Each sheet is one PHC, titled "CHC <chc> PHC <phc>" (e.g.
# "CHC Abhanpur PHC Khorpa "  — note some titles have trailing whitespace).
# The block itself ("ABHANPUR"/"DHARSEEVA") is NOT in the sheet name — it's
# read per-row from column 1, same as before. This regex only pulls the
# CHC/PHC names out of the title, used purely for the parent-facility link.
SHEET_TITLE_RE = re.compile(r"^CHC\s+(.+?)\s+PHC\s+(.+?)\s*$", re.IGNORECASE)

# ── SHC name resolution ───────────────────────────────────────────────────
# 23 of 37 SHC names in the roster match a health_facilities SHC name
# exactly once common suffixes ("SHC "/"SSK ") and bracketed qualifiers
# ("(RAIPUR)", "(RPR)") are stripped. The remaining 14 do not match exactly;
# LIKELY_MATCH is my best guess for those, based on name similarity, and is
# NOT confirmed by the client. Villages under these SHCs are still seeded
# (match_status='likely') so registration isn't blocked, but this table
# should be reviewed before you treat it as ground truth — see the
# accompanying note for which ones I'm least confident about.
LIKELY_MATCH = {
    "SHC Chhapora": "SHC Chapora RPR",
    "SHC Doma": "SHC Doma N",
    "SHC Dondekala": "SHC Dandekala",
    "SHC Kurra(Churiya)": "SHC Urkurra",
    "SHC Mandhar": "SHC Madhar",
    "SHC Saloni(RAIPUR)": "SHC Saloni Raipur",
    "SHC Saragaon(RPR)": "SHC Paragaon",
    "SHC Sejbahar": "SHC Sejbahar N",
    "SHC Temri(RPR)": "SHC Tekri",
    "SSK Serkhedi": "SHC Serikhedi",
    "SSK Tulsi(RPR)": "SHC Tulsi RPR",
    # Two roster SHCs ("SHC Tekari" in Abhanpur, "SHC Tekari(RPR)" in
    # Dharsiwa) both resolve to the SAME single facility "SHC Tekari N".
    # Only one of them can be right, or the facility master is missing an
    # entry. I've left both mapped so neither block is blocked, but this
    # pair specifically needs client confirmation before you trust it.
    "SHC Tekari": "SHC Tekari N",
    "SHC Tekari(RPR)": "SHC Tekari N",
}

# No plausible match at all — seeded as unresolved (shc_id = NULL).
NO_MATCH = {"SHC Hasda(RAIPUR)"}

# Vacancy / non-name text seen in the roster's staff-name columns.
VACANT_PATTERNS = re.compile(
    r"^\s*(-{1,3}|vecant|vacant|not\s+sac?tioned.*|n/?a|ss)\s*$", re.IGNORECASE
)


def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\(.*?\)", "", (s or "").lower())
                  .replace("shc", "").replace("ssk", ""))


def clean_worker(raw_name, raw_mobile):
    """Return (name, mobile, status) — never store vacancy text as a name."""
    name = (str(raw_name).strip() if raw_name is not None else "")
    mobile = (str(raw_mobile).strip() if raw_mobile is not None else "") or None

    if not name or VACANT_PATTERNS.match(name):
        if name and "not sac" in name.lower():
            return None, None, WorkerStatus.not_sanctioned
        return None, None, WorkerStatus.vacant

    # Phone sanity check: Indian mobiles are 10 digits (ignoring +91/spaces).
    digits = re.sub(r"\D", "", mobile or "")
    if mobile and len(digits) not in (10, 12):  # 12 = with country code
        return name, mobile, WorkerStatus.data_issue

    return name, mobile, WorkerStatus.active


def parse_roster(filepath: str):
    wb = load_workbook(filepath, read_only=True)
    villages = []           # dicts
    workers = {}             # (shc_roster_name, role) -> (name, mobile, status)
    shc_to_phc = {}          # shc_roster_name -> phc_roster_name (from row data)
    phc_to_chc_sheet = {}    # phc_roster_name -> chc_name (from the sheet title)
    unmatched_sheets = []

    for sheet_name in wb.sheetnames:
        title_match = SHEET_TITLE_RE.match(sheet_name.strip())
        if not title_match:
            unmatched_sheets.append(sheet_name)
            continue
        chc_from_title = title_match.group(1).strip()

        ws = wb[sheet_name]
        rows_in_sheet = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row[0]:
                continue
            rows_in_sheet += 1
            # Columns per earlier inspection:
            # 0 idx, 1 block, 2 phc, 3 shc, 4 village(+code), 5 cho, 6 cho_mobile,
            # 7 rho_f, 8 rho_f_mobile, 9 rho_m, 10 rho_m_mobile, 11 anm2, 12 anm2_mobile
            block = (row[1] or "").strip()
            phc_name = (row[2] or "").strip()
            shc_name = (row[3] or "").strip()
            village_raw = (row[4] or "").strip()

            m = re.match(r"^(.*?)\s*\(([\w*]+)\)\s*$", village_raw)
            if m:
                vname, vcode = m.group(1).strip(), m.group(2)
            else:
                vname, vcode = village_raw, None

            villages.append({
                "name": vname,
                "code": vcode,
                "block": block,
                "shc_roster_name": shc_name,
            })

            if shc_name and phc_name:
                shc_to_phc[shc_name] = phc_name
            if phc_name:
                phc_to_chc_sheet[phc_name] = chc_from_title

            for role, name_col, mob_col in [
                (WorkerRole.cho, 5, 6),
                (WorkerRole.rho_female, 7, 8),
                (WorkerRole.rho_male, 9, 10),
                (WorkerRole.anm_2nd, 11, 12),
            ]:
                key = (shc_name, role)
                if key not in workers:
                    workers[key] = clean_worker(
                        row[name_col] if len(row) > name_col else None,
                        row[mob_col] if len(row) > mob_col else None,
                    )

        if rows_in_sheet == 0:
            unmatched_sheets.append(f"{sheet_name} (title matched, but 0 data rows found)")

    if unmatched_sheets:
        print(f"WARNING: sheets not parsed (title didn't match 'CHC ... PHC ...'): {unmatched_sheets}")

    return villages, workers, shc_to_phc, phc_to_chc_sheet


async def resolve_shc_ids(db):
    """Build {normalized_roster_shc_name: (facility_id, match_status)}."""
    result = await db.execute(select(HealthFacility).where(HealthFacility.facility_type == "SHC"))
    by_norm = {norm(f.name): f.id for f in result.scalars().all()}

    mapping = {}
    all_roster_shc_names = set()

    def resolve(shc_roster_name: str):
        if shc_roster_name in all_roster_shc_names:
            return mapping[shc_roster_name]
        all_roster_shc_names.add(shc_roster_name)

        if shc_roster_name in NO_MATCH:
            mapping[shc_roster_name] = (None, "unresolved")
            return mapping[shc_roster_name]

        exact_id = by_norm.get(norm(shc_roster_name))
        if exact_id:
            mapping[shc_roster_name] = (exact_id, "exact")
            return mapping[shc_roster_name]

        likely_name = LIKELY_MATCH.get(shc_roster_name)
        if likely_name:
            likely_id = by_norm.get(norm(likely_name))
            if likely_id:
                mapping[shc_roster_name] = (likely_id, "likely")
                return mapping[shc_roster_name]

        mapping[shc_roster_name] = (None, "unresolved")
        return mapping[shc_roster_name]

    return resolve


async def seed():
    print("Reading roster Excel...")
    villages, workers, shc_to_phc, phc_to_chc_sheet = parse_roster(ROSTER_FILE)
    print(f"Parsed {len(villages)} village rows, {len(workers)} worker slots.")

    async with AsyncSessionLocal() as db:
        resolve = await resolve_shc_ids(db)

        unresolved_villages = 0
        likely_villages = 0

        for v in villages:
            shc_id, status = resolve(v["shc_roster_name"])
            if status == "unresolved":
                unresolved_villages += 1
            elif status == "likely":
                likely_villages += 1

            existing = None
            if v["code"]:
                res = await db.execute(select(Village).where(Village.code == v["code"]))
                existing = res.scalar_one_or_none()
            else:
                res = await db.execute(
                    select(Village).where(Village.name == v["name"], Village.block == v["block"])
                )
                existing = res.scalar_one_or_none()

            if existing:
                existing.name = v["name"]
                existing.block = v["block"]
                existing.shc_id = shc_id
                existing.match_status = status
            else:
                db.add(Village(
                    name=v["name"], code=v["code"], block=v["block"],
                    district="Raipur", shc_id=shc_id, match_status=status,
                ))

        await db.commit()
        print(f"Villages seeded. unresolved={unresolved_villages}, likely={likely_villages}")

        # ── Workers ──────────────────────────────────────────────────────
        # Two different roster SHC names can resolve to the SAME real
        # facility (this happened for the Tekari pair — see LIKELY_MATCH
        # comment above). That means two different worker records could
        # target the same (facility_id, role) slot, which the DB's unique
        # constraint correctly rejects. Resolve + dedupe in memory first,
        # keeping the first-seen entry and reporting every collision, so a
        # bad roster mapping surfaces as a printed warning, not a crash.
        skipped_no_facility = 0
        resolved_workers = {}   # (facility_id, role) -> (name, mobile, status, source_roster_name)
        collisions = []

        for (shc_roster_name, role), (name, mobile, status) in workers.items():
            shc_id, _ = resolve(shc_roster_name)
            if not shc_id:
                skipped_no_facility += 1
                continue
            key = (shc_id, role)
            if key in resolved_workers:
                collisions.append((shc_roster_name, resolved_workers[key][3], role))
                continue
            resolved_workers[key] = (name, mobile, status, shc_roster_name)

        if collisions:
            print(f"WARNING: {len(collisions)} worker rows skipped due to two roster SHC "
                  f"names resolving to the same facility (kept the first one seen):")
            for lost_name, kept_name, role in collisions:
                print(f"  - role={role}: '{lost_name}' data discarded in favour of '{kept_name}'")

        # One bulk fetch of existing rows for the facilities involved, so
        # we never rely on a select-after-add seeing a not-yet-flushed row.
        facility_ids = {fid for fid, _ in resolved_workers}
        existing_res = await db.execute(
            select(HealthWorker).where(HealthWorker.facility_id.in_(facility_ids))
        )
        existing_by_key = {(w.facility_id, w.role): w for w in existing_res.scalars().all()}

        for (facility_id, role), (name, mobile, status, _src) in resolved_workers.items():
            existing_worker = existing_by_key.get((facility_id, role))
            if existing_worker:
                existing_worker.name = name
                existing_worker.mobile = mobile
                existing_worker.status = status
            else:
                db.add(HealthWorker(facility_id=facility_id, role=role, name=name, mobile=mobile, status=status))

        await db.commit()
        print(f"Workers seeded. skipped (no resolved SHC)={skipped_no_facility}, "
              f"skipped (facility collision)={len(collisions)}")

        # ── PHC -> CHC parent links (from each sheet's title: "CHC X PHC Y") ──
        # Direct from the source, not guessed from block name.
        chc_res = await db.execute(select(HealthFacility).where(HealthFacility.facility_type == "CHC"))
        chc_by_norm = {norm(f.name): f for f in chc_res.scalars().all()}

        phc_res = await db.execute(select(HealthFacility).where(HealthFacility.facility_type == "PHC"))
        phc_by_norm = {norm(f.name): f for f in phc_res.scalars().all()}

        phc_linked = 0
        phc_not_found = []
        chc_not_found = []
        for phc_roster_name, chc_from_title in phc_to_chc_sheet.items():
            phc = phc_by_norm.get(norm(phc_roster_name))
            chc = chc_by_norm.get(norm(f"CHC {chc_from_title}"))
            if not phc:
                phc_not_found.append(phc_roster_name)
                continue
            if not chc:
                chc_not_found.append(chc_from_title)
                continue
            if phc.parent_facility_id != chc.id:
                phc.parent_facility_id = chc.id
                phc_linked += 1

        # ── SHC -> PHC parent links (only for the 5 PHCs the roster covers) ──
        shc_linked = 0
        shc_res_all = await db.execute(select(HealthFacility).where(HealthFacility.facility_type == "SHC"))
        shc_by_id = {f.id: f for f in shc_res_all.scalars().all()}

        for shc_roster_name, phc_roster_name in shc_to_phc.items():
            shc_id, _ = resolve(shc_roster_name)
            if not shc_id or not phc_roster_name:
                continue
            phc = phc_by_norm.get(norm(phc_roster_name))
            if not phc:
                continue
            shc_obj = shc_by_id.get(shc_id)
            if shc_obj and shc_obj.parent_facility_id != phc.id:
                shc_obj.parent_facility_id = phc.id
                shc_linked += 1

        await db.commit()
        print(f"Facility hierarchy linked. PHC->CHC={phc_linked}, SHC->PHC={shc_linked}")
        if phc_not_found:
            print(f"  PHC names from roster not found in health_facilities: {phc_not_found}")
        if chc_not_found:
            print(f"  CHC names (from sheet titles) not found in health_facilities: {chc_not_found}")

        print("\nFlagged for client/admin review (match_status != 'exact'):")
        res = await db.execute(select(Village).where(Village.match_status != "exact"))
        for v in res.scalars().all():
            print(f"  - {v.name} (block={v.block}, status={v.match_status}, shc_id={v.shc_id})")


if __name__ == "__main__":
    asyncio.run(seed())